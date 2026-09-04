# -*- coding: utf-8 -*-
"""Materialise selected LAB tasks as plain text so subagents can read them."""
import json, os, glob, re, shutil
import docx

ROOT='/home/claude/harvey-labs/tasks'
OUT='/home/claude/exp/tasks'

SEL = [
 ('arbitration-international-dispute-resolution','identify-deficiencies-and-vulnerabilities-in-counterpartys-statement-of-claim','misrep-en'),
 ('arbitration-international-dispute-resolution','analyze-arbitration-award-for-new-york-convention-enforcement-defenses','nyc'),
 ('arbitration-international-dispute-resolution','draft-challenge-to-arbitrator-appointment','arbch'),
 ('arbitration-international-dispute-resolution','identify-issues-in-arbitrator-disclosure-statement','arbch'),
 ('arbitration-international-dispute-resolution','draft-application-to-set-aside-arbitration-award','nyc'),
 ('arbitration-international-dispute-resolution','compare-arbitration-clause-vs-institutional-rules-compliance-audit','model-law'),
]

COV = {
 'misrep-deceit': r'derry v peek|deceit|misrepresent|entire agreement|non-reliance|inducement|reliance on|fraudulent|remoteness|hadley',
 'model-law':     r'model law|justifiable doubts|arbitrator.*(challeng|disclos|impartial|independen)|iba guidelines|article 12|article 13',
 'nyc-enforce':   r'new york convention|article v|set(ting)? aside|enforcement of the award|public policy|article 34',
}

def docx2txt(p):
    try:
        d=docx.Document(p)
        parts=[x.text for x in d.paragraphs if x.text.strip()]
        for t in d.tables:
            for r in t.rows:
                cells=[c.text.strip() for c in r.cells]
                if any(cells): parts.append(' | '.join(cells))
        return '\n'.join(parts)
    except Exception as e:
        return f'[unreadable: {e}]'

manifest=[]
for area,slug,fam in SEL:
    src=os.path.join(ROOT,area,slug)
    if not os.path.isdir(src): print('MISSING',slug); continue
    t=json.load(open(os.path.join(src,'task.json')))
    dst=os.path.join(OUT,slug); os.makedirs(os.path.join(dst,'documents'),exist_ok=True)
    nchars=0
    for f in sorted(glob.glob(os.path.join(src,'documents','*'))):
        base=os.path.basename(f); ext=os.path.splitext(base)[1].lower()
        if ext=='.docx': txt=docx2txt(f)
        elif ext in ('.eml','.txt','.md','.csv'): txt=open(f,encoding='utf-8',errors='replace').read()
        elif ext=='.xlsx':
            import openpyxl
            wb=openpyxl.load_workbook(f,data_only=True); rows=[]
            for ws in wb:
                rows.append(f'## sheet: {ws.title}')
                for r in ws.iter_rows(values_only=True):
                    if any(v is not None for v in r): rows.append(' | '.join('' if v is None else str(v) for v in r))
            txt='\n'.join(rows)
        else: continue
        nchars+=len(txt)
        open(os.path.join(dst,'documents',base+'.txt'),'w',encoding='utf-8').write(txt)
    # covered legal-standard criteria
    cov=[]
    for c in t['criteria']:
        x=c['title']+' '+c['match_criteria']
        fams=[k for k,p in COV.items() if re.search(p,x,re.I)]
        if fams: cov.append({'id':c['id'],'title':c['title'],'match_criteria':c['match_criteria'],'families':fams})
    json.dump({'slug':slug,'family':fam,'title':t['title'],'instructions':t['instructions'],
               'deliverables':t['deliverables'],'n_criteria':len(t['criteria']),
               'covered_criteria':cov}, open(os.path.join(dst,'meta.json'),'w'), ensure_ascii=False, indent=1)
    manifest.append({'slug':slug,'family':fam,'docs':len(glob.glob(os.path.join(dst,'documents','*'))),
                     'doc_chars':nchars,'n_criteria':len(t['criteria']),'n_covered':len(cov)})
json.dump(manifest,open('/home/claude/exp/manifest.json','w'),indent=1)
print(f"{'covered':>8} {'total':>6} {'docs':>5} {'chars':>8}  task")
for m in manifest: print(f"{m['n_covered']:8d} {m['n_criteria']:6d} {m['docs']:5d} {m['doc_chars']:8d}  {m['slug'][:52]}")
print('TOTAL covered criteria:', sum(m['n_covered'] for m in manifest))
